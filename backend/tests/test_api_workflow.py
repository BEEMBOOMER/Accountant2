from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import main
from backend.app.services import BankingService
from backend.app.workbook import WorkbookStore


ROOT = Path(__file__).resolve().parents[2]


def api(tmp_path, monkeypatch):
    service = BankingService(WorkbookStore(tmp_path / "banking_master.xlsx"))
    monkeypatch.setattr(main, "service", service)
    return TestClient(main.app), service


def upload(client, endpoint, filename, bank=None, account_id=None):
    data = {}
    if bank:
        data["bank"] = bank
    if account_id:
        data["account_id"] = account_id
    content = (ROOT / "LearningMaterials" / filename).read_bytes()
    return client.post(
        endpoint,
        data=data,
        files={"file": (filename, content, "text/csv")},
    )


def test_preview_import_deduplication_and_export(tmp_path, monkeypatch):
    client, service = api(tmp_path, monkeypatch)

    preview = upload(client, "/api/imports/preview", "CSVData.csv", account_id="commbank-main")
    assert preview.status_code == 200
    assert preview.json()["inserted"] == 16
    assert service.state()[0] == []

    detected = upload(client, "/api/imports/detect", "CSVData.csv")
    assert detected.status_code == 200
    assert detected.json() == {"bank": "CommBank", "account_required": True}

    first = upload(client, "/api/imports", "CSVData.csv", account_id="commbank-main")
    assert first.status_code == 200
    assert first.json()["inserted"] == 16
    assert first.json()["latest_month"] == "2026-08"

    repeated = upload(client, "/api/imports", "CSVData.csv", account_id="commbank-main")
    assert repeated.status_code == 200
    assert repeated.json()["inserted"] == 0
    assert repeated.json()["duplicates"] == 16

    westpac = upload(client, "/api/imports", "Data_export_11082026.csv", "Westpac")
    assert westpac.status_code == 200
    assert westpac.json()["inserted"] == 43

    months = client.get("/api/analytics/months")
    assert months.status_code == 200
    assert months.json() == {"months": ["2026-08", "2026-07"], "latest": "2026-08"}
    latest_analytics = client.get("/api/analytics").json()
    assert latest_analytics["month"] == "2026-08"

    transactions = client.get("/api/transactions", params={"month": "2026-08"})
    assert transactions.status_code == 200
    assert transactions.json()["total"] > 0
    assert all("id" in item and "category_id" in item for item in transactions.json()["transactions"])
    pending = client.get("/api/transactions", params={"review_required": "true"}).json()
    assert pending["total"] > 0
    assert all(item["review_required"] for item in pending["transactions"])

    exported = client.get("/api/export")
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.content), read_only=True)
    assert workbook.sheetnames == [
        "Transactions", "Categories", "Rules", "Accounts", "Import History", "Monthly Summary"
    ]
    assert workbook["Transactions"].max_row == 60
    workbook.close()


def test_import_auto_detects_bank_and_rejects_unknown_structure(tmp_path, monkeypatch):
    client, _ = api(tmp_path, monkeypatch)
    westpac = upload(client, "/api/imports", "Data_export_11082026.csv")
    assert westpac.status_code == 200
    assert westpac.json()["bank"] == "Westpac"
    unknown = client.post(
        "/api/imports/detect",
        files={"file": ("unknown.csv", b"merchant,value\nshop,10\n", "text/csv")},
    )
    assert unknown.status_code == 400
    assert "not recognized" in unknown.json()["detail"]


def test_transfer_confirmation_updates_both_legs_and_savings(tmp_path, monkeypatch):
    client, _ = api(tmp_path, monkeypatch)
    upload(client, "/api/imports", "CSVData.csv", "CommBank", "commbank-main")
    upload(client, "/api/imports", "Data_export_11082026.csv", "Westpac")

    before = client.get("/api/analytics", params={"month": "2026-08"}).json()
    rows = client.get("/api/transactions", params={"month": "2026-08", "transaction_type": "transfer"}).json()["transactions"]
    suspected = next(item for item in rows if item["transfer_status"] == "suspected")

    response = client.patch(f"/api/transactions/{suspected['id']}/transfer", json={"status": "confirmed"})
    assert response.status_code == 200

    after_rows = client.get("/api/transactions", params={"month": "2026-08", "transaction_type": "transfer"}).json()["transactions"]
    confirmed_group = [item for item in after_rows if item["transfer_group_id"] == suspected["transfer_group_id"]]
    assert len(confirmed_group) == 2
    assert all(item["transfer_status"] == "confirmed" and item["internal_transfer"] for item in confirmed_group)

    after = client.get("/api/analytics", params={"month": "2026-08"}).json()
    assert after["income"] <= before["income"]
    assert after["expenditure"] <= before["expenditure"]
    if any(item["amount"] > 0 and item["account"] == "734001872217" for item in confirmed_group):
        assert after["savings"] > 0


def test_category_edit_rule_and_required_replacement(tmp_path, monkeypatch):
    client, _ = api(tmp_path, monkeypatch)
    upload(client, "/api/imports", "CSVData.csv", "CommBank", "commbank-main")
    categories = client.get("/api/categories").json()
    groceries = next(item for item in categories if item["name"] == "Groceries")
    miscellaneous = next(item for item in categories if item["name"] == "Miscellaneous")
    tx = client.get("/api/transactions").json()["transactions"][0]

    changed = client.patch(
        f"/api/transactions/{tx['id']}/category",
        json={"category_id": groceries["id"], "create_exact_rule": True},
    )
    assert changed.status_code == 200
    assert changed.json()["category"] == "Groceries"
    assert any(rule["kind"] == "merchant" for rule in client.get("/api/rules").json())

    blocked = client.delete(f"/api/categories/{groceries['id']}")
    assert blocked.status_code == 400
    replaced = client.delete(
        f"/api/categories/{groceries['id']}",
        params={"replacement_category_id": miscellaneous["id"]},
    )
    assert replaced.status_code == 200
    assert changed.json()["id"] in {item["id"] for item in client.get("/api/transactions").json()["transactions"]}


def test_reimbursements_have_their_own_api_type_and_income_filter_excludes_them(tmp_path, monkeypatch):
    client, _ = api(tmp_path, monkeypatch)
    response = client.post(
        "/api/imports",
        json={
            "account_id": "commbank-main",
            "source_file": "wifi.csv",
            "csv_text": "01/08/2026,13.33,Fast Transfer From DHANUK DE SILVA Dhanuk Wifi,113.34",
        },
    )
    assert response.status_code == 200

    rows = client.get("/api/transactions").json()["transactions"]
    assert rows[0]["type"] == "reimbursement"
    assert rows[0]["expense_reimbursement"] is True
    assert client.get("/api/transactions", params={"transaction_type": "income"}).json()["total"] == 0
    assert client.get("/api/transactions", params={"transaction_type": "reimbursement"}).json()["total"] == 1
