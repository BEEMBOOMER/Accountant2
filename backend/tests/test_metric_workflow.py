"""Exercise dashboard drilldowns and corrections through the public API."""

from decimal import Decimal

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import main
from backend.app.services import BankingService
from backend.app.workbook import WorkbookStore


def test_dashboard_correction_survives_excel_reload_and_duplicate_import(tmp_path, monkeypatch):
    workbook = tmp_path / "banking_master.xlsx"
    service = BankingService(WorkbookStore(workbook))
    monkeypatch.setattr(main, "service", service)
    client = TestClient(main.app)
    statement = (
        '12/09/2026,-100.00,"PayTo Webull Securities (Australia) Pty monoova-webull-example",900.00\n'
        '13/09/2026,-300.00,"Transfer To WestPac Savings CommBank App Savings",600.00\n'
        '14/09/2026,-50.00,"Unknown housing payment",550.00\n'
    )

    def upload():
        return client.post(
            "/api/imports",
            data={"bank": "CommBank", "account_id": "commbank-main"},
            files={"file": ("statement.csv", statement.encode(), "text/csv")},
        )

    assert upload().status_code == 200
    for metric, expected in (("investments", "100"), ("savings", "300"), ("expenditure", "50")):
        response = client.get("/api/transactions", params={"month": "2026-09", "metric": metric})
        assert response.status_code == 200
        rows = response.json()["transactions"]
        assert len(rows) == 1
        summary = client.get("/api/analytics", params={"month": "2026-09"}).json()
        assert Decimal(str(summary[metric])) == Decimal(expected)

    savings_row = client.get("/api/transactions", params={"month": "2026-09", "metric": "savings"}).json()["transactions"][0]
    housing = next(c for c in client.get("/api/categories").json() if c["name"] == "Housing")
    correction = client.patch(
        f"/api/transactions/{savings_row['id']}/classification",
        json={"classification": "expense", "category_id": housing["id"]},
    )
    assert correction.status_code == 200, correction.text
    assert correction.json()["category"] == "Housing"
    assert client.get("/api/transactions", params={"month": "2026-09", "metric": "savings"}).json()["total"] == 0
    assert client.get("/api/transactions", params={"month": "2026-09", "metric": "expenditure"}).json()["total"] == 2

    reopened = BankingService(WorkbookStore(workbook))
    monkeypatch.setattr(main, "service", reopened)
    repeated = upload()
    assert repeated.status_code == 200
    assert repeated.json()["inserted"] == 0
    assert repeated.json()["duplicates"] == 3
    final_rows = client.get("/api/transactions", params={"month": "2026-09", "metric": "expenditure"}).json()["transactions"]
    corrected = next(row for row in final_rows if row["id"] == savings_row["id"])
    assert corrected["category"] == "Housing"
    assert corrected["original_description"] == savings_row["original_description"]
    assert corrected["fingerprint"] == savings_row["fingerprint"]
    summary = client.get("/api/analytics", params={"month": "2026-09"}).json()
    assert Decimal(str(summary["expenditure"])) == Decimal("350")
    assert Decimal(str(summary["savings"])) == Decimal("0")
    assert Decimal(str(summary["investments"])) == Decimal("100")
    assert client.get("/api/transactions", params={"month": "2026-08", "metric": "investments"}).json()["total"] == 0


def test_invalid_corrections_leave_workbook_unchanged(tmp_path, monkeypatch):
    service = BankingService(WorkbookStore(tmp_path / "banking_master.xlsx"))
    monkeypatch.setattr(main, "service", service)
    service.import_csv(
        '12/09/2026,-50.00,"Housing payment",950.00\n',
        account="commbank-main", source_file="statement.csv",
    )
    client = TestClient(main.app)
    transaction_id = service.state()[0][0].transaction_id
    before = service.store.path.read_bytes()
    for body in (
        {"classification": "expense", "category_id": "nonexistent-category"},
        {"classification": "income"},
        {"classification": "reimbursement"},
        {"classification": "unknown"},
    ):
        response = client.patch(f"/api/transactions/{transaction_id}/classification", json=body)
        assert response.status_code == 400, response.text
        assert service.store.path.read_bytes() == before


def test_destination_evidence_prevents_equal_amount_savings_double_count(tmp_path):
    service = BankingService(WorkbookStore(tmp_path / "banking_master.xlsx"))
    service.import_csv(
        'Bank Account,Date,Narrative,Debit Amount,Credit Amount,Balance\n'
        '034001844999,06/08/2026,DEPOSIT-OSKO PAYMENT 2427258 ANAY VENUGOPAL Savings account,,2000.00,2000.00\n',
        account=None, source_file="westpac.csv",
    )
    service.import_csv(
        '04/08/2026,-2000.00,"Transfer to xx2990 CommBank app",100.00\n'
        '06/08/2026,2000.00,"Transfer from xx2990 CommBank app",2100.00\n'
        '06/08/2026,-2000.00,"Transfer To WestPac Savings CommBank App Savings account",100.00\n',
        account="commbank-main", source_file="commbank.csv",
    )
    summary = service.analytics("2026-08")
    assert summary["savings"] == Decimal("2000")
    assert summary["income"] == Decimal("0")
    assert summary["expenditure"] == Decimal("0")
    rows = service.metric_transactions("savings", "2026-08")
    assert len(rows) == 1
    assert rows[0].bank == "Westpac"
    service.set_classification(rows[0].transaction_id, "transfer")
    assert service.analytics("2026-08")["savings"] == Decimal("0")
    service.apply_financial_patterns()
    assert service.metric_transactions("savings", "2026-08") == []


def test_manual_outgoing_savings_survives_recalculation(tmp_path):
    service = BankingService(WorkbookStore(tmp_path / "banking_master.xlsx"))
    service.import_csv('12/09/2026,-200.00,"My deposit",800.00\n', account="commbank-main", source_file="statement.csv")
    transaction = service.state()[0][0]
    service.set_classification(transaction.transaction_id, "savings")
    service.apply_financial_patterns()
    assert service.analytics("2026-09")["savings"] == Decimal("200")
    assert len(service.metric_transactions("savings", "2026-09")) == 1
    assert service.state()[0][0].savings_transfer is True


def test_later_imported_savings_receipt_across_month_boundary_counts_once(tmp_path):
    service = BankingService(WorkbookStore(tmp_path / "banking_master.xlsx"))
    service.import_csv(
        '31/08/2026,-300.00,"Transfer To WestPac Savings CommBank App Savings",700.00\n',
        account="commbank-main", source_file="outgoing.csv",
    )
    assert service.analytics("2026-08")["savings"] == Decimal("300")
    service.import_csv(
        'Bank Account,Date,Narrative,Debit Amount,Credit Amount,Balance\n'
        '734001872217,01/09/2026,DEPOSIT-OSKO PAYMENT ANAY VENUGOPAL Savings account,,300.00,300.00\n',
        account=None, source_file="incoming.csv",
    )
    assert service.analytics("2026-08")["savings"] == Decimal("0")
    assert service.analytics("2026-09")["savings"] == Decimal("300")
    assert service.metric_transactions("savings", "2026-08") == []
    assert len(service.metric_transactions("savings", "2026-09")) == 1


def test_legacy_upgrade_preserves_confirmed_choices_and_writes_once(tmp_path, monkeypatch):
    service = BankingService(WorkbookStore(tmp_path / "banking_master.xlsx"))
    service.import_csv('12/09/2026,-200.00,"My deposit",800.00\n', account="commbank-main", source_file="statement.csv")
    transaction = service.state()[0][0]
    service.set_classification(transaction.transaction_id, "savings")
    workbook = load_workbook(service.store.path)
    sheet = workbook["Transactions"]
    # Reproduce the pre-upgrade schema and its confirmed decision.
    for column in range(sheet.max_column, 0, -1):
        if sheet.cell(1, column).value in {"classification", "classification_manual", "category_manual", "investment_transfer"}:
            sheet.delete_cols(column)
    headers = [cell.value for cell in sheet[1]]
    sheet.cell(2, headers.index("transfer_group_id") + 1, "legacy-confirmed")
    workbook.save(service.store.path)
    workbook.close()
    saves = []
    save = service.store._atomic_save

    def counted_save(workbook):
        saves.append(True)
        save(workbook)

    monkeypatch.setattr(service.store, "_atomic_save", counted_save)
    assert service.upgrade_classification_schema()["schema_changed"] is True
    assert len(saves) == 1
    row = service.state()[0][0]
    assert row.classification_manual is True
    assert row.classification == "savings"
    assert row.original_description == transaction.original_description
    before = service.store.path.read_bytes()
    assert service.upgrade_classification_schema()["schema_changed"] is False
    assert service.store.path.read_bytes() == before
    assert len(saves) == 1
