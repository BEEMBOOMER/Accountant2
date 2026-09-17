from __future__ import annotations

import os
import shutil
import tempfile
import threading
import uuid
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from .category_lookup import CURATED_CATEGORIES, CURATED_RULES
from .models import Account, Category, ImportHistory, MonthlySummary, Rule, Transaction


TX_HEADERS = [
    "transaction_id", "fingerprint", "date", "bank", "account", "original_description",
    "normalized_description", "amount", "balance", "transaction_type", "category", "category_manual",
    "transfer_status", "internal_transfer", "transfer_group_id", "savings_transfer",
    "classification", "classification_manual", "investment_transfer", "expense_reimbursement",
    "review_required", "source_file", "source_row", "imported_at",
]
ACCOUNT_HEADERS = ["id", "bank", "name", "identifier", "is_savings", "currency"]
CATEGORY_HEADERS = ["id", "name", "enabled"]
RULE_HEADERS = ["id", "match_type", "pattern", "category", "enabled"]
HISTORY_HEADERS = ["id", "bank", "account", "source_file", "imported_at", "processed", "inserted", "duplicates", "suspected_transfers", "review_required", "rejected", "errors"]
SUMMARY_HEADERS = ["month", "income", "expenditure", "net_cashflow", "savings", "investments", "category", "category_spend"]

DEFAULT_CATEGORIES = CURATED_CATEGORIES
DEFAULT_RULES = CURATED_RULES
DEFAULT_ACCOUNTS = [
    ("commbank-main", "CommBank", "CommBank Everyday", "xx2990", False),
    ("westpac-034001844999", "Westpac", "Westpac Choice", "034001844999", False),
    ("westpac-734001872217", "Westpac", "Westpac Life", "734001872217", True),
]


def _decimal(value: Any, default: Decimal | None = None) -> Decimal | None:
    if value is None or value == "":
        return default
    try:
        return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"invalid monetary value {value!r}") from exc


def _bool(value: Any) -> bool:
    return value is True or str(value).strip().lower() in {"true", "1", "yes"}


class WorkbookStore:
    """The Excel workbook is the sole persistence layer for application state."""

    def __init__(self, path: str | Path | None = None):
        root = Path(__file__).resolve().parents[2]
        self.path = Path(path) if path else root / "data" / "banking_master.xlsx"
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._lock = threading.RLock()
        self.ensure()

    def ensure(self) -> None:
        with self._lock:
            if self.path.exists():
                try:
                    wb = load_workbook(self.path, read_only=True)
                    missing = set(("Transactions", "Categories", "Rules", "Accounts", "Import History", "Monthly Summary")) - set(wb.sheetnames)
                    wb.close()
                    if missing:
                        raise RuntimeError(f"missing sheets: {', '.join(sorted(missing))}")
                    return
                except Exception as exc:
                    raise RuntimeError(f"unable to read workbook {self.path}; it was left unchanged: {exc}") from exc
            wb = Workbook()
            wb.remove(wb.active)
            for name, headers in (("Transactions", TX_HEADERS), ("Categories", CATEGORY_HEADERS),
                                  ("Rules", RULE_HEADERS), ("Accounts", ACCOUNT_HEADERS),
                                  ("Import History", HISTORY_HEADERS), ("Monthly Summary", SUMMARY_HEADERS)):
                ws = wb.create_sheet(name)
                ws.append(headers)
                ws.freeze_panes = "A2"
            for name in DEFAULT_CATEGORIES:
                wb["Categories"].append([str(uuid.uuid4()), name, True])
            for match_type, pattern, category in DEFAULT_RULES:
                wb["Rules"].append([str(uuid.uuid4()), match_type, pattern, category, True])
            for account_id, bank, name, identifier, savings in DEFAULT_ACCOUNTS:
                wb["Accounts"].append([account_id, bank, name, identifier, savings, "AUD"])
            self._atomic_save(wb)

    def install_classification_schema(self) -> bool:
        """Safely add classification columns to an existing workbook.

        This explicit, idempotent migration is intentionally never invoked by
        construction or reads, so opening the supplied financial workbook does
        not modify it.  The next normal write also uses the current schema.
        """
        with self._lock:
            wb = self.load()
            changed = False
            try:
                for sheet, headers in (("Transactions", TX_HEADERS), ("Monthly Summary", SUMMARY_HEADERS)):
                    ws = wb[sheet]
                    present = {str(cell.value) for cell in ws[1] if cell.value is not None}
                    for header in headers:
                        if header not in present:
                            ws.cell(row=1, column=ws.max_column + 1, value=header)
                            present.add(header)
                            changed = True
                if changed:
                    self._atomic_save(wb)
                return changed
            finally:
                wb.close()

    def has_classification_schema(self) -> bool:
        """Read-only check used to protect pre-upgrade manual transfer choices."""
        with self._lock:
            wb = self.load()
            try:
                return "classification_manual" in {str(cell.value) for cell in wb["Transactions"][1] if cell.value is not None}
            finally:
                wb.close()

    def _atomic_save(self, wb) -> None:
        fd, temp_name = tempfile.mkstemp(prefix="banking_master_", suffix=".xlsx", dir=self.path.parent)
        os.close(fd)
        temp_path = Path(temp_name)
        try:
            wb.save(temp_path)
            check = load_workbook(temp_path, read_only=True)
            check.close()
            os.replace(temp_path, self.path)
        finally:
            if temp_path.exists():
                temp_path.unlink()

    def load(self):
        self.ensure()
        try:
            return load_workbook(self.path)
        except Exception as exc:
            raise RuntimeError(f"unable to read workbook {self.path}: {exc}") from exc

    def save(self, wb) -> None:
        with self._lock:
            self._atomic_save(wb)

    @staticmethod
    def _rows(ws):
        headers = [str(x.value) for x in next(ws.iter_rows(min_row=1, max_row=1))]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(value is not None for value in row):
                continue
            yield dict(zip(headers, row))

    def read_all(self) -> tuple[list[Transaction], list[Category], list[Rule], list[Account], list[ImportHistory], list[MonthlySummary]]:
        with self._lock:
            wb = self.load()
            try:
                txs = []
                legacy_classification = "classification_manual" not in {
                    str(cell.value) for cell in wb["Transactions"][1] if cell.value is not None
                }
                for row in self._rows(wb["Transactions"]):
                    try:
                        transaction = Transaction(
                            transaction_id=str(row.get("transaction_id", "")), fingerprint=str(row.get("fingerprint", "")),
                            date=row.get("date") if isinstance(row.get("date"), date) else date.fromisoformat(str(row.get("date"))),
                            bank=str(row.get("bank", "")), account=str(row.get("account", "")),
                            original_description=str(row.get("original_description", "")), normalized_description=str(row.get("normalized_description", "")),
                            amount=_decimal(row.get("amount")), balance=_decimal(row.get("balance")),
                            transaction_type=str(row.get("transaction_type", "")), category=str(row.get("category", "Miscellaneous")), category_manual=_bool(row.get("category_manual")),
                            transfer_status=str(row.get("transfer_status", "none")), internal_transfer=_bool(row.get("internal_transfer")),
                            transfer_group_id=row.get("transfer_group_id"), savings_transfer=_bool(row.get("savings_transfer")),
                            classification=str(row.get("classification")) if row.get("classification") else None,
                            classification_manual=_bool(row.get("classification_manual")),
                            investment_transfer=_bool(row.get("investment_transfer")),
                            expense_reimbursement=_bool(row.get("expense_reimbursement")),
                            review_required=_bool(row.get("review_required")), source_file=str(row.get("source_file", "")),
                            source_row=int(row.get("source_row") or 0), imported_at=row.get("imported_at") if isinstance(row.get("imported_at"), datetime) else datetime.fromisoformat(str(row.get("imported_at"))),
                        )
                        # Old workbooks had no durable manual-decision field.
                        # Preserve explicit grouped confirmations/rejections on
                        # their first upgraded write instead of re-detecting
                        # them as automatic matches.
                        if legacy_classification and (transaction.transfer_status == "rejected" or (
                            transaction.transfer_status == "confirmed"
                            and transaction.transfer_group_id
                            and not str(transaction.transfer_group_id).startswith("automatic:")
                        )):
                            transaction.classification_manual = True
                            if transaction.transfer_status == "confirmed":
                                transaction.classification = "savings" if transaction.savings_transfer else "transfer"
                            else:
                                transaction.classification = "reimbursement" if transaction.expense_reimbursement else ("income" if transaction.amount > 0 else "expense")
                        txs.append(transaction)
                    except Exception as exc:
                        row_number = txs.__len__() + 2
                        raise RuntimeError(f"invalid Transactions row near row {row_number}: {exc}") from exc
                categories = [Category(id=str(r.get("id")), name=str(r.get("name")), enabled=_bool(r.get("enabled"))) for r in self._rows(wb["Categories"])]
                rules = [Rule(id=str(r.get("id")), match_type=str(r.get("match_type")), pattern=str(r.get("pattern")), category=str(r.get("category")), enabled=_bool(r.get("enabled"))) for r in self._rows(wb["Rules"])]
                accounts = [Account(id=str(r.get("id")), bank=str(r.get("bank")), name=str(r.get("name")), identifier=str(r.get("identifier")), is_savings=_bool(r.get("is_savings")), currency=str(r.get("currency") or "AUD")) for r in self._rows(wb["Accounts"])]
                history = [ImportHistory(id=str(r.get("id")), bank=str(r.get("bank")), account=str(r.get("account")), source_file=str(r.get("source_file")), imported_at=r.get("imported_at") if isinstance(r.get("imported_at"), datetime) else datetime.fromisoformat(str(r.get("imported_at"))), processed=int(r.get("processed") or 0), inserted=int(r.get("inserted") or 0), duplicates=int(r.get("duplicates") or 0), suspected_transfers=int(r.get("suspected_transfers") or 0), review_required=int(r.get("review_required") or 0), rejected=int(r.get("rejected") or 0), errors=str(r.get("errors") or "").split("\n") if r.get("errors") else []) for r in self._rows(wb["Import History"])]
                summary_rows: dict[str, MonthlySummary] = {}
                for r in self._rows(wb["Monthly Summary"]):
                    month = str(r.get("month"))
                    summary = summary_rows.setdefault(
                        month,
                        MonthlySummary(
                            month=month,
                            income=_decimal(r.get("income"), Decimal("0")),
                            expenditure=_decimal(r.get("expenditure"), Decimal("0")),
                            net_cashflow=_decimal(r.get("net_cashflow"), Decimal("0")),
                            savings=_decimal(r.get("savings"), Decimal("0")),
                            investments=_decimal(r.get("investments"), Decimal("0")),
                        ),
                    )
                    category = r.get("category")
                    if category:
                        summary.category_totals[str(category)] = _decimal(r.get("category_spend"), Decimal("0")) or Decimal("0")
                summaries = list(summary_rows.values())
                return txs, categories, rules, accounts, history, summaries
            finally:
                wb.close()

    def replace_all(self, transactions, categories, rules, accounts, history, summaries) -> None:
        with self._lock:
            wb = self.load()
            try:
                for name, headers in (("Transactions", TX_HEADERS), ("Categories", CATEGORY_HEADERS), ("Rules", RULE_HEADERS), ("Accounts", ACCOUNT_HEADERS), ("Import History", HISTORY_HEADERS), ("Monthly Summary", SUMMARY_HEADERS)):
                    ws = wb[name]
                    ws.delete_rows(1, ws.max_row)
                    ws.append(headers)
                ws = wb["Transactions"]
                for t in transactions:
                    ws.append([t.transaction_id, t.fingerprint, t.date, t.bank, t.account, t.original_description, t.normalized_description, t.amount, t.balance, t.transaction_type, t.category, t.category_manual, t.transfer_status, t.internal_transfer, t.transfer_group_id, t.savings_transfer, t.classification, t.classification_manual, t.investment_transfer, t.expense_reimbursement, t.review_required, t.source_file, t.source_row, t.imported_at.isoformat()])
                ws = wb["Categories"]
                for c in categories: ws.append([c.id, c.name, c.enabled])
                ws = wb["Rules"]
                for r in rules: ws.append([r.id, r.match_type, r.pattern, r.category, r.enabled])
                ws = wb["Accounts"]
                for a in accounts: ws.append([a.id, a.bank, a.name, a.identifier, a.is_savings, a.currency])
                ws = wb["Import History"]
                for h in history: ws.append([h.id, h.bank, h.account, h.source_file, h.imported_at.isoformat(), h.processed, h.inserted, h.duplicates, h.suspected_transfers, h.review_required, h.rejected, "\n".join(h.errors)])
                ws = wb["Monthly Summary"]
                for s in summaries:
                    category_rows = sorted(s.category_totals.items()) or [(None, None)]
                    for category, spend in category_rows:
                        ws.append([s.month, s.income, s.expenditure, s.net_cashflow, s.savings, s.investments, category, spend])
                self._atomic_save(wb)
            finally:
                wb.close()
